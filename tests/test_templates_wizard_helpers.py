#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tester för wizard-hjälparna i script/templates.py + roundtrip SV/EN."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from script.templates import (  # noqa: E402
    build_template_body,
    create_template,
    find_placeholders_in_text,
    list_placeholders,
    parse_subject_and_body,
    sanitize_template_filename,
)


class TestFindPlaceholdersInText(unittest.TestCase):
    def test_all_three_formats(self) -> None:
        text = "Hej ::modell::, {{reg}} och ((ruta))!"
        self.assertEqual(
            find_placeholders_in_text(text), ["modell", "reg", "ruta"]
        )

    def test_dedupe_case_insensitive_order_kept(self) -> None:
        text = "::MODELL:: x ::modell:: y ::Reg:: z ::REG::"
        self.assertEqual(find_placeholders_in_text(text), ["MODELL", "Reg"])

    def test_empty_and_no_placeholders(self) -> None:
        self.assertEqual(find_placeholders_in_text(""), [])
        self.assertEqual(find_placeholders_in_text("Vanlig text."), [])


class TestBuildTemplateBody(unittest.TestCase):
    def test_swedish_only(self) -> None:
        self.assertEqual(build_template_body("Hej!", ""), "Hej!")

    def test_with_english_block(self) -> None:
        body = build_template_body("Hej!", "Hello!")
        self.assertEqual(body, "Hej!\n\n[en]\nHello!\n[/en]")

    def test_roundtrip_through_parser(self) -> None:
        body = build_template_body(
            "Modell: ::modell::", "Model: ::modell::\nVIN: ::vin::"
        )
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "mall.txt"
            create_template(path, "Ämne ::modell::", body)
            subject_sv, sv = parse_subject_and_body(
                path, {"modell": "Volvo"}, include_english=False,
                keep_unknown=True,
            )
            _subject_en, en = parse_subject_and_body(
                path, {"modell": "Volvo"}, include_english=True,
                keep_unknown=True,
            )
            self.assertEqual(subject_sv, "Ämne Volvo")
            self.assertIn("Modell: Volvo", sv)
            self.assertNotIn("VIN", sv)
            self.assertIn("Model: Volvo", en)
            self.assertIn("::vin::", en)
            self.assertEqual(
                list_placeholders(path), ["modell", "vin"]
            )

    def test_example_template_in_repo(self) -> None:
        example = ROOT / "Mallar" / "Exempel – Offertförfrågan Bilglas.txt"
        self.assertTrue(example.exists(), "Exempelmallen saknas i Mallar/")
        subject, sv = parse_subject_and_body(
            example,
            {"modell": "Saab 9-3", "reg": "ABC 123", "ruta": "Vindruta"},
            include_english=False, keep_unknown=True,
        )
        _s, en = parse_subject_and_body(
            example,
            {"modell": "Saab 9-3", "reg": "ABC 123", "ruta": "Vindruta"},
            include_english=True, keep_unknown=True,
        )
        self.assertIn("Saab 9-3", subject)
        self.assertIn("Bilmodell: Saab 9-3", sv)
        self.assertNotIn("VIN:", sv)
        self.assertIn("Car model: Saab 9-3", en)
        self.assertIn("VIN: ::vin::", en)


class TestSanitizeTemplateFilename(unittest.TestCase):
    def test_basic(self) -> None:
        self.assertEqual(
            sanitize_template_filename("Offertförfrågan Bilglas"),
            "Offertförfrågan Bilglas.txt",
        )

    def test_illegal_windows_chars_removed(self) -> None:
        self.assertEqual(
            sanitize_template_filename('Min "offert" <Bilglas>: v2?'),
            "Min offert Bilglas v2.txt",
        )
        self.assertEqual(
            sanitize_template_filename("a/b\\c|d"), "abcd.txt"
        )

    def test_extension_forced(self) -> None:
        self.assertEqual(
            sanitize_template_filename("Mall.txt", ".docx"), "Mall.docx"
        )
        self.assertEqual(
            sanitize_template_filename("Mall.docx", "html"), "Mall.html"
        )
        self.assertEqual(
            sanitize_template_filename("Mall", ".exe"), "Mall.txt"
        )

    def test_reserved_names_and_trailing_dots(self) -> None:
        self.assertEqual(sanitize_template_filename("CON"), "_CON.txt")
        self.assertEqual(sanitize_template_filename("Mall..."), "Mall.txt")
        self.assertEqual(sanitize_template_filename("  Mellan  slag  "),
                         "Mellan slag.txt")

    def test_empty_results(self) -> None:
        self.assertEqual(sanitize_template_filename(""), "")
        self.assertEqual(sanitize_template_filename("???"), "")
        self.assertEqual(sanitize_template_filename("..."), "")

    def test_swedish_letters_preserved(self) -> None:
        self.assertEqual(
            sanitize_template_filename("Återkoppling övning ä"),
            "Återkoppling övning ä.txt",
        )


class TestCreateFormatsRoundtrip(unittest.TestCase):
    """Wizarden ska kunna skapa TXT + HTML (+ DOCX om python-docx finns)."""

    BODY = build_template_body(
        "Hej ::modell::", "Hello ::modell::"
    )

    def _roundtrip(self, filename: str) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / filename
            create_template(path, "Test ::modell::", self.BODY)
            self.assertTrue(path.exists())
            subject, body = parse_subject_and_body(
                path, {"modell": "XC60"}, keep_unknown=True
            )
            self.assertEqual(subject, "Test XC60")
            self.assertIn("Hej XC60", body)

    def test_txt(self) -> None:
        self._roundtrip("a.txt")

    def test_html(self) -> None:
        self._roundtrip("a.html")

    def test_docx(self) -> None:
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx ej installerat")
        self._roundtrip("a.docx")


if __name__ == "__main__":
    unittest.main()
